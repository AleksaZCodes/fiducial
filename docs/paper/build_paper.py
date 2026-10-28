"""Build the Fiducial paper (.docx) on the IEEESTEC template.

The template is the author's own IEEESTEC 2025 paper (Outreach), converted
from .doc to .docx: its styles carry the conference formatting (two columns,
auto-numbered headings, "Fig. N." captions, "TABLE N." heads, "[N]"
references). Every paragraph after the author block is removed and the text
below is written in its place.

    python3 build_paper.py TEMPLATE.docx FIGURE_DIR OUT.docx
"""

import copy
import sys

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

TEMPLATE, FIGS, OUT = sys.argv[1:4]

# Only the conference line changes; the DOI and "Original scientific paper"
# lines stay as the template has them.
HEADER = ["IEEESTEC – 19th Student Projects Conference, Niš, 2026."]

TITLE = "Fiducial: Keeping Multidisciplinary Products Built with AI Agents Consistent"

AUTHOR = [
    "Aleksa Zdravković",
    "12th grade high school student",
    "ETŠ “Nikola Tesla” Niš",
    "Niš, Serbia",
    "aleksazdravkovic@proton.me",
    "ORCID: 0009-0004-1535-1839",
]

import json as _json
import re as _re
from pathlib import Path as _Path

_HERE = _Path(__file__).resolve().parent
_RESULTS = _json.loads((_HERE / "markers/results.json").read_text())
_SUM = _RESULTS["summary"]
_BY_ID = {r["id"]: r for r in _RESULTS["results"]}
VERSION = _re.search(r'^version\s*=\s*"([^"]+)"',
                     (_HERE.parents[1] / "Cargo.toml").read_text(), _re.M).group(1)

def _count(gate):
    """How many markers a gate caught first."""
    return sum(1 for r in _RESULTS["results"]
               if not r["control"] and r["caught_by"] and r["caught_by"][0] == gate)


_C1 = _BY_ID["C1"]

def _load(name):
    return _json.loads((_HERE / "markers" / name).read_text())


_HELD = _load("results-heldout.json")
_BASE = _load("results-baseline.json")
_HBASE = _load("results-heldout-baseline.json")
_H = _HELD["summary"]["markers_caught"]
_HB = _HBASE["summary"]["markers_caught"]
_B = _BASE["summary"]["markers_caught"]
_HB_W = "none" if _HB == 0 else str(_HB)
# The text names the held-out misses; a rerun that changes them must change it.
assert {r["id"] for r in _HELD["results"] if not r["pass"]} == {"N5", "N6", "N16", "N19", "N20"}, \
    "the held-out misses changed: update FINDINGS"

ALL_CAUGHT = _SUM["markers_caught"] == _SUM["markers"]
CONTROLS_OK = _SUM["controls_passed"] == _SUM["controls"]

_PIN_S = round(next(g["seconds"] for g in _C1["gates"] if g["gate"] == "derive"))
_PIN_N = len([f for f in _C1["changed"] if "/generated/" in f])
_N = _SUM["markers_caught"]
_ALL = "all 20" if ALL_CAUGHT else f"{_N} of the 20"

ABSTRACT = (
    "This paper presents Fiducial, an open-source platform for building "
    "products that span software and hardware with AI coding agents. Agents "
    "now write each part of a product quickly, but the facts the parts share "
    "are still copied between them and drift apart. Fiducial declares every "
    "fact once, generates every artifact from it – code, translations, "
    "circuit board and enclosure – and checks every change, so the parts "
    "stay consistent with each other; whether a declared fact is itself "
    "correct remains for the person to judge. Of 20 realistic mistakes, the "
    f"checks first caught 15, and {_ALL} once each miss had become a rule. "
    f"On 20 new mistakes written afterwards they caught {_H}, where the "
    f"same project’s ordinary CI caught {_HB_W}."
)

KEYWORDS = (
    "AI agents, single source of truth, code generation, consistency "
    "checking, hardware–software co-design, continuous integration"
)

# ("h1", text) | ("h2", text) | ("p", text) | ("fig", file[@inches], caption)
# | ("table", caption, header, rows, widths) — a row that is a str is a group
BODY = [
    ("h1", "Introduction"),
    ("p",
     "Building a product used to take a team: someone for the circuit board, "
     "someone for the firmware, others for the app and the enclosure. AI "
     "coding agents such as Claude Code [@claude] and Codex [@codex] are "
     "changing that: a firmware driver, a web page or a script that models a "
     "case can be written in minutes, so one person can now attempt a whole "
     "product."),
    ("p",
     "Building such products – first Outreach [@outreach], an off-grid "
     "messaging system of radios, gateways, an app and printed cases, and "
     "then the platform described here – has shown where the difficulty has "
     "moved. Writing each part is no longer the slow step; keeping the parts "
     "in agreement is. In Outreach, the firmware and the app disagreed on "
     "the layout of a message, and a printed case did not fit its board. "
     "Neither part was wrong on its own: a shared fact had been copied by "
     "hand, one copy changed, and the mistake appeared only when the parts "
     "met (Fig. 1a)."),
    ("fig", "fig1-one-fact.png",
     "One fact used in four places: (a) copied by hand, one copy changes and "
     "nothing notices; (b) declared once, every copy is generated from it and "
     "checked."),
    ("p",
     "Agents speed this problem up rather than remove it. An agent sees only "
     "the codebase it works in, and its output is usually right but never "
     "guaranteed; more code, written faster in more disconnected places, "
     "gives an inconsistent product sooner."),
    ("p",
     "What is missing is one shared and checked context for every "
     "discipline: a single place where a product’s facts live, from which "
     "every part is produced, and which refuses a change that makes the parts "
     "disagree (Fig. 1b). This paper presents Fiducial [@fiducial], an "
     "open-source platform built to provide it, compares it with existing "
     "approaches and with ordinary CI, and tests its checks on mistakes they "
     "were not built for."),

    ("h1", "Existing Approaches"),
    ("p",
     "Several kinds of tool address parts of this problem; each keeps one "
     "kind of fact, or one pair of disciplines, in agreement, but not the "
     "product as a whole (Table I)."),
    ("p",
     "Model-driven and model-based systems engineering keep a design in one "
     "central model and generate or check other views from it [@schmidt], "
     "[@brambilla], in languages such as SysML [@sysml], with methods for "
     "managing inconsistencies between views [@nuseibeh], [@herzig]. They "
     "are the closest in spirit, but the model lives in a specialised "
     "graphical tool, code is still written by hand beside it, and an agent "
     "cannot easily work in it. Engineering suites instead link two tools, "
     "such as a circuit board and its mechanical model [@altium]; the "
     "firmware, the app and the translations stay outside."),
    ("p",
     "Design-as-code tools write hardware as text – a circuit in atopile "
     "[@atopile], a case in build123d [@build123d] – which makes it writable "
     "by an agent, but each covers one discipline: Fiducial’s first design "
     "planned to use atopile and found that the case, the layout and the "
     "cost would still copy the board’s facts. Interface languages such as "
     "Protocol Buffers [@protobuf] declare a message once for both sides, "
     "which removes the drift met in Outreach, but only for messages."),
    ("table", "Existing Approaches and Where They Stop",
     ("Approach", "Keeps in agreement", "Stops at"), [
        ("Model-based engineering", "views of one model",
         "graphical tool; code by hand"),
        ("ECAD–MCAD links", "board and case", "two disciplines, two tools"),
        ("Design as code", "one discipline", "facts shared with others"),
        ("Interface languages", "messages, both sides", "everything else"),
        ("Build systems, CI", "outputs with files", "two files that disagree"),
        ("Agent context files", "nothing; they advise", "not enforced"),
        ("Fiducial", "every declared fact", "facts not yet declared"),
    ], (1500, 1550, 1500)),
    ("p",
     "Build systems such as Bazel [@bazel] and Nix [@nix] rebuild exactly "
     "what changed, identifying results by fingerprints of their inputs. "
     "They make sure outputs follow from files, not that files agree with "
     "each other: the same fact typed into two files builds without "
     "complaint, in them and in a CI pipeline built on them."),
    ("p",
     "Coding agents follow context files such as AGENTS.md [@agentsmd], "
     "which describe a repository’s rules but do not enforce them, and the "
     "code they write is often plausible but wrong. Stricter tests find "
     "errors in programs that benchmark tests [@chen] had passed [@liu], "
     "about 40 % of the programs in one security study were vulnerable "
     "[@pearce], and most real repository issues remain unsolved "
     "[@jimenez]. Models also cannot reliably correct their own reasoning "
     "without an outside signal [@huang]. A known answer is to pair a "
     "generator that can be wrong with a checker that cannot "
     "[@kambhampati], but no existing checker spans a whole product."),
    ("p",
     "Fiducial keeps what works in each – text that agents can write, "
     "outputs tied to inputs by fingerprints, established engines for hard "
     "problems, and checks that decide – and changes the unit. It tracks "
     "facts rather than files or tools, across every discipline of one "
     "product, so any fact – a message, a pin, a colour, a size – has "
     "exactly one home."),

    ("h1", "The Approach"),
    ("p",
     "Fiducial treats a product as facts and artifacts. Facts are the "
     "decisions that define it: which pin a sensor uses, what a button says "
     "in each language, the size of the board, the layout of a message. "
     "Artifacts are everything built from them: code, schematics, boards, "
     "translations, enclosure models. Drift is when an artifact stops "
     "agreeing with a fact, or with another artifact that uses it "
     "(Fig. 1a). Three rules keep drift out (Fig. 1b)."),
    ("p",
     "Declare once. Every fact is written in exactly one place, a "
     "declaration: a typed text file in the repository, with a structure a "
     "program can check. A fact that exists once has no second copy to "
     "disagree with."),
    ("p",
     "Derive everything else. Every artifact that depends on a fact is "
     "generated from its declaration by one command and never edited by "
     "hand. Generation is deterministic, so the product is reproducible and "
     "any difference is meaningful."),
    ("p",
     "Check every change. Before a change is accepted, checks confirm that "
     "every artifact still matches what its declarations produce and that "
     "the declarations are valid; a failure names the declaration and the "
     "file at fault."),
    ("p",
     "The rules check agreement, not truth. A fact declared wrongly – the "
     "wrong sensor address, a resistor of the wrong value – is copied "
     "faithfully into every artifact, and the product is then consistently "
     "wrong. Some such values break a rule and are refused, such as two "
     "colours with too little contrast or a resistor that puts a voltage "
     "outside its window, but whether a declaration is right is still "
     "decided by the person, the datasheet and the bench."),
    ("p",
     "These rules change the role of an AI agent: it still edits "
     "declarations and code, but it is no longer trusted to keep anything "
     "consistent. Its mistakes are not prevented, but many are caught before "
     "they reach the product, with a message precise enough for the agent "
     "to correct itself. Hence the name: a fiducial is a fixed reference "
     "point, a universal measure that everything is aligned to and compared "
     "against, so it stays uniform. The declarations are that reference for "
     "a product."),

    ("h1", "Implementation"),
    ("p",
     "Fiducial is a command-line tool, fid, built over six weeks by one "
     "person working with AI agents. Fig. 2 shows how a change moves through "
     "its parts."),
    ("h2", "One Structure for Every Product"),
    ("p",
     "fid new creates a product repository with a fixed layout, and fid add "
     "enables capabilities. A capability packages one discipline: its "
     "declarations, the generator that turns them into artifacts, its checks "
     "and instructions for agents. There are 23, from web applications and "
     "translations to firmware and hardware. Because every product has the "
     "same structure, whoever knows one knows them all, and a new product "
     "starts from parts that already work. fid upgrade brings platform "
     "improvements into each product while keeping its own edits."),
    ("h2", "One Typed Language Underneath"),
    ("p",
     "The platform, and the code products share through it, is written in "
     "Rust, a compiled language with a strict type system. The choice follows "
     "from the approach: every mistake the compiler refuses is one the checks "
     "never have to catch, and it is refused before anything runs, with a "
     "message naming the file and line – the signal an agent needs. Rust also "
     "needs no runtime or garbage collector, so one body of code can run "
     "wherever a product does. Its core is written without the standard "
     "library: the same message format, identifiers and typed physical "
     "quantities compile into a microcontroller’s firmware, into the fid "
     "tool, and to WebAssembly, which runs them in the browser and on edge "
     "servers. The TypeScript types the web code uses are generated from the "
     "Rust ones, never written by hand, so the boundary between the two "
     "languages is not kept in step by hand either. TypeScript is used for interfaces, and "
     "Python only at build time, to drive the engines."),
    ("h2", "Declarations and Generation"),
    ("p",
     "Declarations are plain text files in version control, from the "
     "strings of each language to every part of a board with the net each "
     "pin connects to. Text was chosen over a graphical modelling tool "
     "because coding agents read and write it well, and every change then "
     "goes through ordinary code review. fid "
     "derive runs each capability’s generator, which reads the declarations "
     "and writes the artifacts, from a typed module of translations to a "
     "firmware pin map or a KiCad circuit board [@kicad]. A lock file records a "
     "cryptographic hash – a short fingerprint of a file’s contents – for "
     "every input and output, so the platform knows exactly what each "
     "artifact was made from."),
    ("h2", "Checks"),
    ("p",
     "Checks run in continuous integration, the automatic build that runs on "
     "every proposed change, and a change is merged only when they pass. "
     "There are three kinds. Freshness: fid derive --check compares every "
     "recorded fingerprint with the repository and fails if a declaration "
     "changed without its artifacts being regenerated, or if a generated "
     "file was edited by hand. Rules: generation refuses declarations that "
     "break a rule, such as a language missing a string, two colours with too "
     "little contrast, parts that cost more than the declared limit, or a "
     "board that does not fit its case. Consumers: compilers and type "
     "checkers read the generated files, so code that still uses a renamed "
     "fact stops compiling."),
    ("h2", "Engines for Hard Problems"),
    ("p",
     "Some generation steps are real engineering problems: where to place "
     "the parts on a board, how to connect them, how to build a case around "
     "them. The first version placed parts with a hand-written search, and "
     "each new product added special rules until they began to interact. Now "
     "the platform only writes the problem down and hands it to an "
     "established engine. A constraint solver, CP-SAT [@cpsat], is given the "
     "rules a placement must satisfy and finds one that does, or proves none "
     "exists; build123d [@build123d] builds the case; KiCad routes and checks the "
     "board. Every placement rule is one of eight general kinds, such as "
     "inside, apart or near, so a new product needs new declarations rather "
     "than new code."),
    ("fig", "fig2-loop.png",
     "How a change moves through Fiducial: declarations are edited, "
     "generation produces the artifacts with the help of established "
     "engines, and the checks either accept the change or return a message "
     "naming the declaration at fault."),
    ("fig", "fig3-sensor-stick.png",
     "The sensor stick, generated from its declaration: (a) the closed case; "
     "(b) the routed board inside it, with the sensor under its vent."),
    ("h2", "Agents in the Loop"),
    ("p",
     "Every product carries an AGENTS.md file and one instruction file per "
     "capability, so any agent starts from the same description of the "
     "product, and a hook stops it from editing a generated file. Agents fix "
     "a problem where the fact lives, never in the output. The platform grows "
     "the same way: a mistake that gets past the checks becomes a general "
     "rule for every product."),

    ("h1", "Example: A Sensor Stick"),
    ("p",
     "The repository contains a complete example: a USB-C stick that "
     "measures temperature and humidity and shows them on a web page "
     "(Fig. 3), with a printed case, a circuit board, Rust firmware and a "
     "browser page. The person writes the product declaration, an outline of "
     "its shape, the firmware’s logic and the page; the routed board with "
     "its schematic and bill of materials, the case, the manufacturing files "
     "and the code both sides use to exchange readings are generated. Moving "
     "the sensor to another pin is one changed line: the check fails until "
     "the product is regenerated, after which the board, the firmware’s pin "
     "map and the page’s types have all followed with no other edit, and a "
     "pin the chip’s sensor bus cannot use stops the firmware compiling."),

    ("h1", "Evaluation"),
    ("p",
     "Checks are only useful if they catch the mistakes that people and "
     "agents actually make. The test borrows the idea of mutation testing "
     "[@demillo]: a realistic mistake is introduced on purpose, the checks "
     "are run, and the result shows whether they noticed. Each mistake was "
     "applied to a clean copy of one of two products – a fresh web product "
     "in English and Serbian, and the sensor stick – the checks ran as in "
     "continuous integration, and the copy was reset. Two sets were used. "
     "The first 20 mistakes cover eight kinds of drift met while building "
     "Fiducial: stale outputs, hand-edited generated files, incomplete or "
     "conflicting declarations, broken rules, code using an old name, values "
     "typed past their declaration, and text left untranslated. They were "
     "also used to improve the platform. The second 20 are held out: written "
     "after release 0.9.11 was fixed, committed before their first run, and "
     "reported without any change made for them. As a baseline, both sets "
     "also ran through ordinary CI alone – the same compilers, type checkers "
     "and tests, with Fiducial’s own two steps removed. Three correct changes "
     "and two runs with no change tested the opposite direction: that the "
     "checks do not fail without a reason. Table II gives the results."),
    ("table", "Injected Mistakes Caught",
     ("Mistakes", "Ordinary CI", "Fiducial 0.9.4", f"Fiducial {VERSION}"), "TABLE1",
     (1660, 1000, 1000, 1000)),
    ("p", "FINDINGS_A"),
    ("p", "FINDINGS_B"),
    ("p",
     "Both sets, and the script that runs them, are in the repository. Both "
     "were written by the author of the platform, so even the "
     "held-out one may favour what it was built to catch; a set written by "
     "others would be stronger. Only two products have been built, neither "
     "yet manufactured and tested, and the test measures the checks, not "
     "whether agents make fewer mistakes with Fiducial than without it."),
    ("h1", "Conclusion"),
    ("p",
     "AI agents have made writing each part of a product fast, and that has "
     "moved the hard problem to keeping the parts in agreement. Better "
     "models will move the line between what a person checks and what is "
     "left to agents, but a probabilistic system cannot guarantee its "
     "output, and even a deterministic tool is only right about what it was "
     "told. Fiducial answers this at the level of the product rather than "
     "the codebase: it keeps the parts consistent with each other, while "
     "whether each declared fact is true remains the person’s to judge. On "
     "the mistakes it was improved against, its checks went from 15 to "
     f"{_N} of 20; on 20 new ones they caught {_H}, where ordinary CI caught "
     f"{_HB_W}, each with a message that names the cause. The next steps "
     "are to test it on mistakes written by others; to explore "
     "retrieval-augmented generation, so that agents are given the "
     "declarations and documentation a task needs rather than whole "
     "repositories; to refine the platform by applying it to new products; "
     "and to measure and lower the tokens agents spend, reasoning tokens "
     "above all. Fiducial and its example are open source [@fiducial]."),
]

# Table II: each set, with and without Fiducial. 0.9.4 ran only the first
# set (15 of 20, every control passed); the rest is read from the results.
def _ctl(res):
    c = [r for r in res["results"] if r["control"]]
    return f"{sum(r['pass'] for r in c)}/{len(c)}"


TABLE1 = [
    ("First set, used to improve", f"{_B}/20", "15/20", f"{_N}/20"),
    ("Held-out set", f"{_HB}/20", "–", f"{_H}/20"),
    ("Correct changes passed", _ctl(_BASE), "5/5", _ctl(_RESULTS)),
]

FINDINGS_A = (
    "On release 0.9.4 the checks caught 15 of the first 20. Each "
    "miss was a fact the platform did not yet own: two declarations edited "
    "without regenerating went unnoticed because only the hardware "
    "generator recorded which files it read; in two, the firmware wrote a "
    "pin number or a sensor address directly; and a Serbian string left in "
    "English passed every structural check. Each became a general rule, and "
    f"release {VERSION} catches "
    + ("all 20" if ALL_CAUGHT else f"{_N} of the 20")
    + ", each with a message naming the declaration at fault. That "
    "result is on mistakes the platform was improved against: it shows the "
    "rules work, not that they generalise."
)
FINDINGS_B = (
    f"The held-out set measures that. The checks caught {_H} of the 20 new "
    "mistakes. They missed a page declared twice, a post dated a month that "
    "does not exist, a board made thicker, a sensor wait typed as a number "
    "and a page comparing humidity to its band by hand; the last two have "
    "rules, but only for narrower patterns. Ordinary CI caught "
    + ("one" if _B == 1 else str(_B))
    + " of the first 20 – a firmware test noticed a swapped message layout – "
    f"and {_HB_W} of the held-out set, because nothing in it compares one "
    "artifact with another. Every correct change passed under both."
)

REFS = {
    "claude": "Claude Code documentation. Available at: "
              "https://docs.anthropic.com/en/docs/claude-code",
    "codex": "OpenAI Codex. Available at: https://openai.com/codex/",
    "outreach": "A. Zdravković, Đ. Jovanović and P. Šubarević, “Outreach: Modular "
                "AI-enhanced off-grid mesh communication system for remote "
                "environments,” IEEESTEC – 18th Student Projects Conference, Niš, 2025.",
    "fiducial": f"A. Zdravković, Fiducial, release {VERSION}, 2026. Available at: "
                "https://github.com/AleksaZCodes/fiducial",
    "schmidt": "D. C. Schmidt, “Guest editor’s introduction: Model-driven "
               "engineering,” Computer, vol. 39, no. 2, pp. 25–31, 2006.",
    "brambilla": "M. Brambilla, J. Cabot and M. Wimmer, Model-Driven Software "
                 "Engineering in Practice, 2nd ed. Morgan & Claypool, 2017.",
    "sysml": "S. Friedenthal, A. Moore and R. Steiner, A Practical Guide to "
             "SysML, 3rd ed. Morgan Kaufmann, 2014.",
    "nuseibeh": "B. Nuseibeh, S. Easterbrook and A. Russo, “Leveraging "
                "inconsistency in software development,” Computer, vol. 33, "
                "no. 4, pp. 24–29, 2000.",
    "herzig": "S. J. I. Herzig, A. Qamar and C. J. J. Paredis, “An approach to "
              "identifying inconsistencies in model-based systems engineering,” "
              "Procedia Computer Science, vol. 28, pp. 354–362, 2014.",
    "altium": "Altium, MCAD CoDesigner, ECAD–MCAD collaboration. Available at: "
              "https://www.altium.com/",
    "atopile": "atopile, a language for designing circuit boards as code. "
               "Available at: https://github.com/atopile/atopile",
    "build123d": "build123d, a Python CAD library on the OpenCascade kernel. "
                 "Available at: https://github.com/gumyr/build123d",
    "protobuf": "Protocol Buffers. Available at: https://protobuf.dev/",
    "bazel": "Bazel, a build and test tool. Available at: https://bazel.build/",
    "nix": "E. Dolstra, M. de Jonge and E. Visser, “Nix: A safe and policy-free "
           "system for software deployment,” in Proc. LISA, 2004, pp. 79–92.",
    "agentsmd": "AGENTS.md, a format for guiding coding agents. Available at: "
                "https://agents.md/",
    "chen": "M. Chen et al., “Evaluating large language models trained on "
            "code,” arXiv:2107.03374, 2021.",
    "liu": "J. Liu, C. S. Xia, Y. Wang and L. Zhang, “Is your code generated by "
           "ChatGPT really correct? Rigorous evaluation of large language models "
           "for code generation,” in Proc. NeurIPS, 2023.",
    "pearce": "H. Pearce, B. Ahmad, B. Tan, B. Dolan-Gavitt and R. Karri, “Asleep "
              "at the keyboard? Assessing the security of GitHub Copilot’s code "
              "contributions,” in Proc. IEEE S&P, 2022, pp. 754–768.",
    "jimenez": "C. E. Jimenez et al., “SWE-bench: Can language models resolve "
               "real-world GitHub issues?,” in Proc. ICLR, 2024.",
    "huang": "J. Huang et al., “Large language models cannot self-correct "
             "reasoning yet,” in Proc. ICLR, 2024.",
    "kambhampati": "S. Kambhampati et al., “Position: LLMs can’t plan, but can "
                   "help planning in LLM-Modulo frameworks,” in Proc. ICML, "
                   "PMLR 235, 2024.",
    "kicad": "KiCad EDA. Available at: https://www.kicad.org/",
    "cpsat": "L. Perron and F. Didier, CP-SAT, Google OR-Tools. Available at: "
             "https://developers.google.com/optimization/cp/cp_solver",
    "demillo": "R. A. DeMillo, R. J. Lipton and F. G. Sayward, “Hints on test data "
               "selection: Help for the practicing programmer,” Computer, vol. 11, "
               "no. 4, pp. 34–41, 1978.",
}

# Number references by first citation, as IEEE style requires: [@key] in the
# text becomes [n], and the list is printed in that order.
_ORDER = []


def _cite(text):
    def num(m):
        if m.group(1) not in _ORDER:
            _ORDER.append(m.group(1))
        return f"[{_ORDER.index(m.group(1)) + 1}]"
    return _re.sub(r"\[@(\w+)\]", num, text)


BODY = [tuple(_cite(x) if isinstance(x, str) and i else x for i, x in enumerate(item))
        for item in BODY]
FINDINGS_A, FINDINGS_B = _cite(FINDINGS_A), _cite(FINDINGS_B)
_unused = set(REFS) - set(_ORDER)
assert not _unused, f"references never cited: {_unused}"
REFERENCES = [REFS[k] for k in _ORDER]

FIG_WIDTH = Inches(3.3)  # the column is 3.38 in; figure labels are 8 pt at this width

d = docx.Document(TEMPLATE)
styles = {s.name.lower(): s for s in d.styles}


def style(name):
    return styles[name.lower()]


def set_text(p, text):
    """Replace a paragraph's text, keeping the first run's formatting."""
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


# The title section came out of the .doc conversion with a 1872-twip right
# margin. Word lays the title out across the full width regardless (the
# Outreach paper's Word PDF centres it on the page), but LibreOffice takes
# the first section's margins as the page's and narrowed every column to
# about 3.05 in, so its preview ran longer than Word's print. Giving the
# title section the body's margins makes both agree with the Word PDF:
# two 3.38 in columns from 0.63 in to 7.64 in.
_BODY_MARGIN = docx.shared.Twips(907)
d.sections[0].left_margin = _BODY_MARGIN
d.sections[0].right_margin = _BODY_MARGIN

# Header: the conference line (first-page header).
for s in d.sections:
    for hf in (s.first_page_header, s.header):
        if hf.is_linked_to_previous:
            continue
        for p in hf.paragraphs:
            if p.text.startswith("IEEESTEC"):
                set_text(p, HEADER[0])

body = d.element.body
children = list(body.iterchildren())

# Title.
title = next(p for p in d.paragraphs if p.style is not None
             and p.style.name == "paper title")
set_text(title, TITLE)

# Author block: one author, one cell, no borders.
tbl = d.tables[0]
tr = tbl.rows[0]._tr
tcs = tr.findall(qn("w:tc"))
for tc in tcs[1:]:
    tr.remove(tc)
grid = tbl._tbl.find(qn("w:tblGrid"))
cols = grid.findall(qn("w:gridCol"))
for c in cols[1:]:
    grid.remove(c)
total = 10336
cols[0].set(qn("w:w"), str(total))
tc = tcs[0]
tcw = tc.find(qn("w:tcPr")).find(qn("w:tcW"))
tcw.set(qn("w:w"), str(total))
borders = tc.find(qn("w:tcPr")).find(qn("w:tcBorders"))
if borders is not None:
    for b in list(borders):
        b.set(qn("w:val"), "nil")
cell_paras = tbl.rows[0].cells[0].paragraphs
for p, text in zip(cell_paras, AUTHOR):
    set_text(p, text)

# Everything from the abstract to the last reference goes.
abstract_el = next(el for el in children if el.tag == qn("w:p")
                   and "Abstract" in "".join(el.itertext()))
start = children.index(abstract_el)
last_sect = max(i for i, el in enumerate(children)
                if el.tag == qn("w:p") and el.find(".//" + qn("w:sectPr")) is not None)
anchor = children[last_sect]
for el in children[start:last_sect]:
    body.remove(el)


def para(style_name, text=None, align=None):
    p = d.add_paragraph(style=style(style_name))
    if text is not None:
        p.add_run(text)
    if align is not None:
        p.alignment = align
    anchor.addprevious(p._p)
    return p


def labelled(style_name, label, text):
    """Abstract—… / Keywords—…: bold-italic label, as in the template."""
    p = para(style_name)
    r = p.add_run(label + "—")
    r.italic = True
    p.add_run(text)
    return p


labelled("Abstract", "Abstract", ABSTRACT)
labelled("Keywords", "Keywords", KEYWORDS)


def keep_with_next(p):
    p.paragraph_format.keep_with_next = True


def add_table(caption, header, rows, widths):
    """A table; a row that is a plain string is a group heading across it.

    Widths are twips; the two-column body is 4860 twips (3.375 in) wide, and
    cell margins take what the widths leave."""
    keep_with_next(para("table head", caption))
    ncols = len(header)
    t = d.add_table(rows=1 + len(rows), cols=ncols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = t._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    tblW.set(qn("w:w"), str(sum(widths)))
    tblW.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))
    for row_i, values in enumerate([header] + list(rows)):
        cells = t.rows[row_i].cells
        if isinstance(values, str):
            cell = cells[0].merge(cells[-1])
            p = cell.paragraphs[0]
            p.style = style("table copy")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(values)
            r.italic = True
            continue
        for col_i, value in enumerate(values):
            cell = cells[col_i]
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(widths[col_i]))
            tcW.set(qn("w:type"), "dxa")
            p = cell.paragraphs[0]
            p.style = style("table col head" if row_i == 0 else "table copy")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            # "--check" must not break at its hyphens.
            p.add_run(value.replace("--", "\u2011\u2011"))
    for row in t.rows[:-1]:
        for cell in row.cells:
            for cp in cell.paragraphs:
                cp.paragraph_format.keep_with_next = True
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "000000")
        b.append(e)
    tblPr.append(b)
    anchor.addprevious(t._tbl)
    # a small gap after the table
    para("Body Text", "")


for item in BODY:
    kind = item[0]
    if kind == "h1":
        keep_with_next(para("Heading 1", item[1]))
    elif kind == "h2":
        keep_with_next(para("Heading 2", item[1]))
    elif kind == "p":
        para("Body Text", {"FINDINGS_A": FINDINGS_A,
                           "FINDINGS_B": FINDINGS_B}.get(item[1], item[1]))
    elif kind == "fig":
        _, name, caption = item
        p = para("Body Text", align=WD_ALIGN_PARAGRAPH.CENTER)
        p.paragraph_format.first_line_indent = 0
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = 0
        p.paragraph_format.space_after = 0
        name, _, inches = name.partition("@")
        width = Inches(float(inches)) if inches else FIG_WIDTH
        p.add_run().add_picture(f"{FIGS}/{name}", width=width)
        keep_with_next(p)
        # The caption ends the figure: it must not pull the heading or
        # paragraph after it along, or a page break leaves a gap.
        para("figure caption", caption).paragraph_format.keep_with_next = False
    elif kind == "table":
        caption, header, rows, widths = item[1:]
        add_table(caption, header, TABLE1 if rows == "TABLE1" else rows, widths)

# References is a component head: Heading 5, unnumbered, as the template says.
keep_with_next(para("Heading 5", "References"))
for ref in REFERENCES:
    para("references", ref)

# The template's own figures are no longer referenced; drop them, or they
# stay inside the file.
xml = d.element.xml
for rid, rel in list(d.part.rels.items()):
    if "image" in rel.reltype and f'"{rid}"' not in xml:
        d.part.drop_rel(rid) if hasattr(d.part, "drop_rel") else d.part.rels.pop(rid)

d.save(OUT)

# Document properties: the template's still say "Paper Title" / "IEEE".
# python-docx would add a second core.xml to this converted package, so the
# existing one is rewritten in place.
import re
import zipfile

with zipfile.ZipFile(OUT) as z:
    parts = {n: z.read(n) for n in z.namelist()}
core = parts["docProps/core.xml"].decode()
for tag, value in (("dc:title", TITLE), ("dc:creator", "Aleksa Zdravković"),
                   ("cp:lastModifiedBy", "Aleksa Zdravković"),
                   ("dc:subject", "IEEESTEC 2026"), ("cp:keywords", KEYWORDS)):
    pattern = rf"<{tag}(\s[^>]*)?>.*?</{tag}>|<{tag}\s*/>"
    element = f"<{tag}>{value}</{tag}>"
    if re.search(pattern, core, flags=re.S):
        core = re.sub(pattern, element, core, flags=re.S)
    else:
        core = core.replace("</cp:coreProperties>", element + "</cp:coreProperties>")
import datetime

now = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
for tag in ("dcterms:created", "dcterms:modified"):
    core = re.sub(rf"(<{tag}[^>]*>).*?(</{tag}>)", rf"\g<1>{now}\g<2>", core)
core = re.sub(r"<cp:revision>\d+</cp:revision>", "<cp:revision>1</cp:revision>", core)
core = re.sub(r"<cp:lastPrinted>.*?</cp:lastPrinted>", "", core)
parts["docProps/core.xml"] = core.encode()

# The .doc → .docx conversion switched automatic hyphenation on. Word did not
# hyphenate the Outreach paper it printed from the same template, and the
# LibreOffice preview cannot (no dictionary), so with it on, Word would break
# lines differently from the preview every layout decision was checked on.
settings = parts["word/settings.xml"].decode()
settings = re.sub(r'<w:autoHyphenation w:val="(true|1)"/>', '', settings)
parts["word/settings.xml"] = settings.encode()
with zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED) as z:
    for name, data in parts.items():
        z.writestr(name, data)
print("wrote", OUT)
